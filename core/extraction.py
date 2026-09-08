import re
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from core.exceptions import ExtractionError

LineMap = list[tuple[str, int | str]]  # (línea limpia, página[PDF] o "Párrafo N"[DOCX])


def _limpiar_linea(texto: str) -> str:
    if not texto:
        return ""
    texto = texto.replace("\\", "")
    texto = texto.replace("\r", " ").replace("\n", " ").replace("\t", " ")
    return re.sub(r"\s+", " ", texto).strip()


def _aplicar_patrones_omitidos(linea: str, patterns: list[re.Pattern]) -> str:
    for pattern in patterns:
        if pattern.search(linea):
            return "[LINEA_OMITIDA]"
    return linea


def _extract_from_pdf(pdf_path: str | Path, patterns: list[re.Pattern]) -> tuple[str, LineMap]:
    texto_completo = ""
    lineas_mapeadas: LineMap = []

    try:
        doc = fitz.open(pdf_path)
    except Exception as e:
        raise ExtractionError(f"No se pudo abrir el PDF '{pdf_path}': {e}") from e

    try:
        for num_pag, page in enumerate(doc, start=1):
            texto_pagina = page.get_text("text")
            for line in texto_pagina.split("\n"):
                linea_limpia = _limpiar_linea(line)
                if not linea_limpia:
                    continue
                linea_limpia = _aplicar_patrones_omitidos(linea_limpia, patterns)
                lineas_mapeadas.append((linea_limpia, num_pag))
                texto_completo += linea_limpia + "\n"
    except Exception as e:
        raise ExtractionError(f"Error al extraer texto de '{pdf_path}': {e}") from e
    finally:
        doc.close()

    return texto_completo, lineas_mapeadas


def _iter_block_items(document: Document):
    """
    Recorre los elementos de nivel superior del cuerpo del documento
    (párrafos y tablas) en el orden real en que aparecen. `Document.paragraphs`
    y `Document.tables` son colecciones separadas que no preservan el orden
    relativo entre ambos tipos, y `.paragraphs` excluye por completo el texto
    dentro de tablas — por eso se itera directamente sobre el XML del body.
    """
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _extract_from_docx(docx_path: str | Path, patterns: list[re.Pattern]) -> tuple[str, LineMap]:
    texto_completo = ""
    lineas_mapeadas: LineMap = []

    try:
        doc = Document(docx_path)
    except Exception as e:
        raise ExtractionError(f"No se pudo abrir el DOCX '{docx_path}': {e}") from e

    def _agregar_linea(texto: str, ubicacion: str) -> None:
        nonlocal texto_completo
        for line in texto.split("\n"):
            linea_limpia = _limpiar_linea(line)
            if not linea_limpia:
                continue
            linea_limpia = _aplicar_patrones_omitidos(linea_limpia, patterns)
            lineas_mapeadas.append((linea_limpia, ubicacion))
            texto_completo += linea_limpia + "\n"

    try:
        idx_parrafo = 0
        idx_tabla = 0
        for block in _iter_block_items(doc):
            if isinstance(block, Paragraph):
                idx_parrafo += 1
                if block.text.strip():
                    _agregar_linea(block.text, f"Párrafo {idx_parrafo}")
                continue

            idx_tabla += 1
            for idx_fila, row in enumerate(block.rows, start=1):
                celdas_vistas: set[int] = set()
                for cell in row.cells:
                    # `row.cells` repite la misma celda por cada columna que
                    # abarca una celda combinada; sin este control esa celda
                    # generaría una línea (y una posible discrepancia)
                    # duplicada por cada columna que ocupa.
                    if id(cell._tc) in celdas_vistas:
                        continue
                    celdas_vistas.add(id(cell._tc))
                    if cell.text.strip():
                        _agregar_linea(cell.text, f"Tabla {idx_tabla}, Fila {idx_fila}")
    except Exception as e:
        raise ExtractionError(f"Error al extraer texto de '{docx_path}': {e}") from e

    return texto_completo, lineas_mapeadas


def extract_text_with_page_mapping(
    file_path: str | Path,
    ignore_line_patterns: list[str] | None = None,
) -> tuple[str, LineMap]:
    """
    Extrae texto plano y un mapeo línea→ubicación (página en PDF, párrafo en DOCX).

    Extracción puramente mecánica: no tiene conocimiento de ningún tipo de
    documento específico. `ignore_line_patterns` permite a quien llame marcar
    líneas de ruido conocidas (p.ej. cabeceras dinámicas de un tipo de
    contrato) como "[LINEA_OMITIDA]" sin que esa lógica viva en el núcleo.
    """
    patterns = [re.compile(p, re.IGNORECASE) for p in (ignore_line_patterns or [])]
    ext = Path(file_path).suffix.lower()

    if ext == ".pdf":
        return _extract_from_pdf(file_path, patterns)
    elif ext == ".docx":
        return _extract_from_docx(file_path, patterns)
    else:
        raise ExtractionError(f"Formato de archivo no soportado: '{ext}'")
